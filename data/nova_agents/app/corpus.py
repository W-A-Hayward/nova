"""Lecture du dossier brut Projet360 (data/starter/Projet360_NOVA_ETUDIANTS).

Chaque fichier du manifeste devient une source. Le texte est extrait selon le format
(courriel, texte, markdown, PDF, Excel, CSV). Les images sont enregistrées sans
transcription: aucun contenu n'est inventé à partir d'une capture.
"""
import csv, email, io, re
from datetime import datetime
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime
from functools import lru_cache
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent / "data" / "starter" / "Projet360_NOVA_ETUDIANTS"
BRIEF = ROOT.parent / "README.txt"  # consignes du défi, hors du dossier de faits

MOIS = {
    "janvier": 1, "janv": 1, "février": 2, "fevrier": 2, "févr": 2, "fevr": 2, "mars": 3,
    "avril": 4, "avr": 4, "mai": 5, "juin": 6, "juillet": 7, "juil": 7, "août": 8, "aout": 8,
    "septembre": 9, "sept": 9, "octobre": 10, "oct": 10, "novembre": 11, "nov": 11,
    "décembre": 12, "decembre": 12, "déc": 12, "dec": 12,
}
_MOIS_ALT = "|".join(sorted(MOIS, key=len, reverse=True))
_DATE_FR = re.compile(rf"\b(\d{{1,2}})\s+({_MOIS_ALT})\s+(20\d\d)\b", re.I)
_DATE_FR_COURTE = re.compile(rf"\b(\d{{1,2}})\s+({_MOIS_ALT})\b", re.I)
_ISO = re.compile(r"\b(20\d\d)-(\d\d)-(\d\d)\b")
_FN_DATE = re.compile(rf"(\d{{2}})({_MOIS_ALT})", re.I)
_TIME = re.compile(rf"^(?:\d{{1,2}}:\d{{2}}\b|\d{{1,2}}\s+(?:{_MOIS_ALT})\b)", re.I)
_BULLET = re.compile(r"^[-•]\s+")
_FIELD = re.compile(r"^[^:\n]{2,60}:\s+\S")
_TICKET = re.compile(r"[A-Z]+-\d+")
_MONEY = re.compile(r"(\d{1,3}(?:[ \u00a0\u202f.]\d{3})+|\d+)\s*\$")


def baseline_au() -> str:
    """Date de référence écrite dans data/starter/README.txt (30 septembre 2026, 09 h, Montréal)."""
    text = BRIEF.read_text(encoding="utf-8")
    m = re.search(r"(\d{1,2}) septembre (\d{4}) à (\d{2}) h(?:\s*(\d{2}))?", text, re.I)
    if not m:
        raise RuntimeError(f"date de référence introuvable dans {BRIEF}")
    day, year, hour, minute = int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4) or 0)
    tz = re.search(r"UTC[−-](\d{2}):(\d{2})", text)
    off = f"-{tz.group(1)}:{tz.group(2)}" if tz else "-04:00"
    return f"{year:04d}-09-{day:02d}T{hour:02d}:{minute:02d}:00{off}"


def _year() -> int:
    return int(baseline_au()[:4])


def _iso(year: int, month: int, day: int) -> str:
    return f"{year:04d}-{month:02d}-{day:02d}"


def _fr_to_iso(day: str, month: str, year: str | None, fallback_year: int | None) -> str | None:
    mo = MOIS.get(month.lower())
    if not mo:
        return None
    y = int(year) if year else fallback_year
    if not y:
        return None
    return _iso(y, mo, int(day))


def filename_date(name: str) -> str | None:
    m = _FN_DATE.search(name)
    if not m:
        return None
    return _fr_to_iso(m.group(1), m.group(2), None, _year())


def header_date(text: str) -> str | None:
    head = text[:900]
    m = _DATE_FR.search(head)
    if m:
        return _fr_to_iso(m.group(1), m.group(2), m.group(3), None)
    m = _ISO.search(head)
    return m.group(0) if m else None


def passage_date(text: str, doc_date: str) -> str:
    """Date du passage: seulement si elle est en tête (commentaire, log, en-tête), sinon date du document."""
    head = text.strip()
    m = _ISO.match(head)
    if m:
        return m.group(0)
    m = _DATE_FR.match(head)
    if m:
        return _fr_to_iso(m.group(1), m.group(2), m.group(3), None) or doc_date
    m = _DATE_FR_COURTE.match(head)
    if m and doc_date and doc_date[:4].isdigit():
        return _fr_to_iso(m.group(1), m.group(2), None, int(doc_date[:4])) or doc_date
    return doc_date or "inconnue"


def _decode_header(value: str | None) -> str:
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return value


def _money(text: str) -> int | None:
    m = _MONEY.search(text.replace("\u00a0", " ").replace("\u202f", " "))
    if not m:
        return None
    return int(re.sub(r"[ \u00a0\u202f.]", "", m.group(1)))


def _pdf_reader(data: bytes):
    from pypdf import PdfReader
    return PdfReader(io.BytesIO(data))


def pdf_text(data: bytes) -> list[tuple[str, str]]:
    reader = _pdf_reader(data)
    out = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append((f"page {i}", text))
    return out


def xlsx_rows(data: bytes) -> list[tuple[str, str]]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if not rows:
                continue
            headers = [str(c).strip() if c not in (None, "") else f"col{j}" for j, c in enumerate(rows[0], 1)]
            for n, row in enumerate(rows[1:], 2):
                pairs, coords = [], []
                for j, (h, c) in enumerate(zip(headers, row), 1):
                    if c in (None, ""):
                        continue
                    val = c.strftime("%Y-%m-%d") if isinstance(c, datetime) else str(c).strip()
                    pairs.append(f"{h}: {val}")
                    coords.append(j)
                if not pairs:
                    continue
                start, end = f"{_col(coords[0])}{n}", f"{_col(coords[-1])}{n}"
                out.append((f"feuille {name}, ligne {n}, cellules {start}:{end}", " | ".join(pairs)))
    finally:
        wb.close()
    return out


def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def csv_rows(text: str) -> list[tuple[str, str]]:
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return []
    headers = rows[0]
    out = []
    for n, row in enumerate(rows[1:], 2):
        pairs = [f"{h}: {v}" for h, v in zip(headers, row) if v.strip()]
        if pairs:
            out.append((f"ligne {n}", " | ".join(pairs)))
    return out


def eml_message(data: bytes) -> email.message.Message:
    return email.message_from_bytes(data)


def eml_view(data: bytes, names: dict[str, str]) -> dict:
    msg = eml_message(data)
    subject = _decode_header(msg.get("subject"))
    author = _decode_header(msg.get("from"))
    author = re.sub(r"\s*<[^>]+>\s*", "", author).strip()
    mid = (msg.get("message-id") or "").strip()
    try:
        sent = parsedate_to_datetime(msg.get("date")).date().isoformat()
    except Exception:
        sent = ""
    bodies, attachments = [], []
    parts = msg.walk() if msg.is_multipart() else [msg]
    for part in parts:
        ctype = part.get_content_type()
        fn = part.get_filename()
        if fn:
            fn = _decode_header(fn)
            other = names.get(fn)
            if other:
                attachments.append(f"Pièce jointe {fn} : même fichier que {other} (ne compte pas comme une confirmation indépendante).")
            else:
                attachments.append(f"Pièce jointe {fn}.")
        elif ctype == "text/plain":
            payload = part.get_payload(decode=True) or b""
            charset = part.get_content_charset() or "utf-8"
            bodies.append(payload.decode(charset, errors="replace").strip())
    body = "\n\n".join(b for b in bodies if b)
    header = f"De: {author}\nDate: {sent}\nObjet: {subject}"
    text = header + "\n\n" + body
    if attachments:
        text += "\n\n" + "\n".join(attachments)
    return {"text": text.strip(), "date": sent, "auteur": author, "message_id": mid, "sujet": subject}


def _atom(line: str) -> bool:
    s = line.strip()
    return bool(_TIME.match(s) or _BULLET.match(s) or _FIELD.match(s))


def split_prose(text: str) -> list[tuple[str, str]]:
    out, buf, buf_start = [], [], 1

    def flush():
        nonlocal buf
        if any(x.strip() for x in buf):
            out.append((f"ligne {buf_start}", "\n".join(x.strip() for x in buf if x.strip())))
        buf = []

    for i, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            flush()
            continue
        if _atom(line):
            flush()
            out.append((f"ligne {i}", line.strip()))
        else:
            if not buf:
                buf_start = i
            buf.append(line)
    flush()
    return out or ([("document", text.strip())] if text.strip() else [])


def _kind(name: str) -> str:
    return Path(name).suffix.lower().lstrip(".")


def read_bytes(name: str, data: bytes, names: dict[str, str] | None = None) -> str:
    """Texte d'un fichier téléversé (même extracteurs que le dossier de départ)."""
    kind = _kind(name)
    names = names or {}
    if kind == "pdf":
        return "\n\n".join(t for _, t in pdf_text(data))
    if kind == "xlsx":
        return "\n".join(t for _, t in xlsx_rows(data))
    if kind == "eml":
        return eml_view(data, names)["text"]
    if kind == "png":
        return f"Capture image {name}. Le contenu visuel n'a pas été transcrit."
    return data.decode("utf-8", errors="replace")


def _autorite(rel: str, name: str, text: str) -> int:
    if "BROUILLON" in name.upper() or re.search(r"(?m)^BROUILLON\b", text):
        return 6
    folder = rel.split("/", 1)[0]
    if folder == "06_Architecture_et_decisions" and name.lower().endswith(".pdf"):
        return 4
    return {
        "01_Courriels": 3,
        "02_Reunions": 2,
        "03_Tickets": 2,
        "04_Documents_projet": 4,
        "05_Contrats_et_finances": 1,
        "06_Architecture_et_decisions": 1,
        "07_Conversations_Teams": 5,
        "08_Archives_et_documents_connexes": 6,
    }.get(folder, 6)


def _pertinence(name: str, text: str) -> str:
    if name in {"README.txt", "MANIFEST.csv"}:
        return "CONSIGNES"
    if name.startswith(("Invitation_", "Newsletter_")):
        return "HORS SUJET"
    if name.startswith("Notes_personnelles") or re.search(r"auteur non identifié|ne vaut rien", text, re.I):
        return "NON FIABLE"
    if "ORION" in name or "INV-778" in name:
        return "HORS SUJET"
    if "preliminaire" in name.lower():
        return "ANCIEN"
    return "NOVA"


def _perime(name: str, text: str, sujet: str) -> str:
    reasons = []
    if "BROUILLON" in name.upper() or re.search(r"\bbrouillon\b", sujet, re.I) or re.search(r"(?m)^BROUILLON\b", text):
        reasons.append("brouillon: ne vaut ni approbation ni engagement")
    if re.search(r"n'est pas mise à jour|non mise à jour|NON mis à jour", text, re.I):
        reasons.append("document non mis à jour après les décisions")
    if re.search(r"préparé avant", text, re.I):
        reasons.append("préparé avant les dernières vérifications")
    if "Architecture_NOVA_v1" in name or re.search(r"version initiale préparée avant", text, re.I):
        reasons.append("version remplacée sur la localisation des données")
    if re.search(r"auteur non identifié|ne vaut rien", text, re.I):
        reasons.append("source non fiable")
    return "; ".join(dict.fromkeys(reasons))


def _auteur(text: str) -> str:
    for pat in (r"Participants\s*:\s*(.+)", r"Demandeur\s*:\s*(.+)", r"Chargée de projet\s*:\s*(.+)"):
        m = re.search(pat, text)
        if m:
            return m.group(1).strip()[:180]
    return ""


def _passages(kind: str, name: str, data: bytes, text: str, names: dict[str, str]) -> tuple[list[tuple[str, str]], dict]:
    extra = {}
    if kind == "eml":
        view = eml_view(data, names)
        extra = view
        return split_prose(view["text"]), extra
    if kind == "pdf":
        pages = pdf_text(data)
        return pages, {"text": "\n\n".join(t for _, t in pages)}
    if kind == "xlsx":
        rows = xlsx_rows(data)
        return rows, {"text": "\n".join(t for _, t in rows)}
    if kind == "csv":
        rows = csv_rows(text)
        return rows, {"text": text}
    if kind == "png":
        ticket = _TICKET.match(Path(name).stem)
        lie = f" Associée au ticket {ticket.group(0)}." if ticket else ""
        note = f"Capture image {name}.{lie} Le contenu visuel n'a pas été transcrit (pas de reconnaissance de texte)."
        return [("image", note)], {"text": note}
    if kind == "md" and len(text) < 2500:
        return [("document", text.strip())], {}
    return split_prose(text), {}


def _load() -> list[dict]:
    files = sorted(p for p in ROOT.rglob("*") if p.is_file())
    names = {}
    for p in files:
        # un même nom en pièce jointe renvoie vers le fichier autonome, pas vers un autre courriel
        if p.suffix.lower() != ".eml":
            names.setdefault(p.name, str(p.relative_to(ROOT)))
    docs = []
    for i, path in enumerate(files, 1):
        rel = str(path.relative_to(ROOT))
        data = path.read_bytes()
        kind = _kind(path.name)
        raw = data.decode("utf-8", errors="replace") if kind not in {"pdf", "xlsx", "png", "eml"} else ""
        passages, extra = _passages(kind, path.name, data, raw, names)
        text = extra.get("text") or raw
        if kind == "xlsx":
            date = filename_date(path.name) or "inconnue"
        else:
            date = extra.get("date") or header_date(text) or filename_date(path.name) or "inconnue"
        stem = Path(path.name).stem
        docs.append({
            "id": f"S{i:02d}",
            "fichier": rel,
            "name": path.name,
            "kind": kind,
            "text": text,
            "date": date,
            "auteur": extra.get("auteur") or _auteur(text),
            "sujet": extra.get("sujet", ""),
            "message_id": extra.get("message_id", ""),
            "autorite": _autorite(rel, path.name, text),
            "pertinence": _pertinence(path.name, text),
            "perime": _perime(path.name, text, extra.get("sujet", "")),
            "brouillon": "BROUILLON" in path.name.upper() or bool(re.search(r"\bbrouillon\b", extra.get("sujet", ""), re.I)),
            "doublon_de": "",
            "ticket_id": stem if rel.startswith("03_Tickets/") and kind == "txt" and re.fullmatch(r"[A-Z]+-\d+", stem) else "",
            "passages": [{"repere": r, "texte": t, "date": passage_date(t, date)} for r, t in passages if t.strip()],
        })
    seen = {}
    for doc in docs:
        mid = doc["message_id"]
        if not mid:
            continue
        if mid in seen:
            doc["doublon_de"] = seen[mid]
            doc["pertinence"] = "NOVA (doublon)"
        else:
            seen[mid] = doc["id"]
    return docs


@lru_cache(maxsize=1)
def documents() -> tuple[dict, ...]:
    if not ROOT.is_dir():
        raise FileNotFoundError(f"dossier source introuvable: {ROOT}")
    return tuple(_load())


def sources() -> list[dict]:
    return [{k: d[k] for k in ("id", "fichier", "date", "auteur", "autorite", "pertinence", "perime", "doublon_de")}
            for d in documents()]


def go_live_cutoff() -> str | None:
    """Date du plus ancien document d'autorité qui approuve le 22 octobre."""
    dates = [d["date"] for d in documents()
             if d["date"] not in ("", "inconnue") and d["autorite"] <= 3
             and re.search(r"22 octobre", d["text"], re.I) and re.search(r"approuv", d["text"], re.I)]
    return min(dates) if dates else None


def _lines(text: str) -> list[str]:
    return [ln.strip() for ln in text.splitlines() if ln.strip()]


def _after(lines: list[str], label: str) -> str:
    for i, ln in enumerate(lines):
        if ln.lower() == label.lower() and i + 1 < len(lines):
            return lines[i + 1]
    return ""


def _invoice(doc: dict) -> dict | None:
    lines = _lines(doc["text"])
    try:
        i = next(n for n, ln in enumerate(lines) if ln.lower() == "montant") + 1
    except StopIteration:
        return None
    items = []
    while i < len(lines) and lines[i].upper() != "TOTAL":
        desc = lines[i]
        i += 1
        if i >= len(lines) or "$" not in lines[i]:
            break
        montant = _money(lines[i])
        i += 1
        if montant is None:
            continue
        ref = ""
        m = re.search(r"\b(CR-\d+)\b", desc)
        if m:
            ref = m.group(1)
        items.append({"desc": desc, "montant": montant, "ref": ref})
    if not items:
        return None
    statut = _after(lines, "Statut")
    statut = "payée" if re.search(r"pay", statut, re.I) else statut.lower()
    return {"id": Path(doc["name"]).stem, "date": _after(lines, "Date") or doc["date"], "statut": statut,
            "source": doc["id"], "lignes": items}


def _change(doc: dict) -> dict | None:
    m = re.search(r"CR-\d+", doc["name"])
    if not m:
        return None
    montant = None
    lines = _lines(doc["text"])
    for i, ln in enumerate(lines):
        if ln.lower().startswith("montant") and i + 1 < len(lines):
            montant = _money(lines[i + 1])
            break
    if montant is None:
        montant = _money(doc["text"])
    approuve = bool(re.search(r"APPROUVÉE", doc["text"])) and "BROUILLON" not in doc["text"].upper()
    date = ""
    mm = _DATE_FR.search(doc["text"])
    if mm:
        date = _fr_to_iso(mm.group(1), mm.group(2), mm.group(3), None) or ""
    return {"id": m.group(0), "montant": montant, "statut": "approuvé" if approuve else "brouillon",
            "date": date or doc["date"], "source": doc["id"]}


def ledger() -> dict:
    """Contrat, changements et factures lus dans les PDF de 05_Contrats_et_finances."""
    contrat = contrat_source = None
    changements, factures = [], []
    for doc in documents():
        if not doc["fichier"].startswith("05_Contrats_et_finances/"):
            continue
        if doc["name"].startswith("CONTRAT"):
            lines = _lines(doc["text"])
            for i, ln in enumerate(lines):
                if ln.lower().startswith("montant maximal") and i + 1 < len(lines):
                    contrat = _money(lines[i + 1])
                    contrat_source = doc["id"]
                    break
        elif doc["name"].startswith("CR-"):
            ch = _change(doc)
            if ch and ch["montant"] is not None:
                changements.append(ch)
        elif doc["name"].startswith("INV-"):
            fac = _invoice(doc)
            if fac:
                factures.append(fac)
    if contrat is None:
        raise RuntimeError("montant contractuel introuvable dans le dossier 05_Contrats_et_finances")
    return {"contrat": contrat, "contrat_source": contrat_source, "changements": changements, "factures": factures}


def ticket_statut(text: str) -> tuple[str, str]:
    m = re.search(r"(?mi)^Statut\s*:\s*(.+)$", text)
    if not m:
        return "", "inconnu"
    raw = m.group(1).strip()
    if re.search(r"en validation", raw, re.I):
        statut = "EN VALIDATION"
    elif re.search(r"ferm", raw, re.I):
        statut = "FERMÉ"
    elif re.search(r"ouvert", raw, re.I):
        statut = "OUVERT"
    else:
        statut = "inconnu"
    return raw, statut
